#-*- coding:utf-8 -*-
from pptx import Presentation

prs = Presentation()#"theme.pptx" 테마설정
f = open('input.txt', encoding='utf-8')

#제목 슬라이드
slide_num = 0
#제목슬라이드는 layout 0번
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
#첫줄은 제목 둘째줄은 부제목
title.text = f.readline()
subtitle.text = f.readline()

#page = [prs.slides[0]]
line_cnt = 1

for line in f:
    #개행문자 입력시 슬라이드 추가
    if(line == '\n'):
        #내용슬라이드는 layout 1번
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        #제목은 미리 받아 삽입
        title_shape.text = f.readline()
        continue

    tf = body_shape.text_frame
    p = tf.add_paragraph()
    p.text = line
    line_cnt = line_cnt + 1

#출력
prs.save('output.pptx')